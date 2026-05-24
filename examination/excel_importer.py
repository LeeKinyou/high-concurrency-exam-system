import json
import logging

from core.constants import Difficulty, QuestionType
from core.exceptions import ValidationError
from core.validators import validate_file_extension, validate_file_size

from .models import Exam, ExamQuestion, Question

logger = logging.getLogger(__name__)

# Excel 列映射：题型, 题干, 选项A-D, 正确答案, 分值, 难度, 解析
EXPECTED_HEADERS = ["题型", "题干", "选项A", "选项B", "选项C", "选项D", "正确答案", "分值", "难度", "解析"]


def import_questions_from_excel(file, exam_id: int | None = None) -> dict:
    """从 Excel 文件导入题目。

    Excel 格式：
    | 题型 | 题干 | 选项A | 选项B | 选项C | 选项D | 正确答案 | 分值 | 难度 | 解析 |
    |------|------|-------|-------|-------|-------|----------|------|------|------|
    | choice | 1+1=? | 1 | 2 | 3 | 4 | B | 10 | easy | |
    | blank | 中国的首都是___ | | | | | 北京 | 10 | easy | |

    Args:
        file: Excel 文件对象
        exam_id: 可选，关联到指定考试

    Returns:
        {"success_count": int, "errors": list[str]}
    """
    validate_file_extension(file)
    validate_file_size(file)

    try:
        from openpyxl import load_workbook

        wb = load_workbook(file, read_only=True)
        ws = wb.active
    except Exception as e:
        raise ValidationError(f"无法读取Excel文件: {e}")

    headers = [str(cell.value).strip() if cell.value else "" for cell in next(ws.iter_rows(max_row=1))]
    col_map = {}
    for i, header in enumerate(headers):
        if header in EXPECTED_HEADERS:
            col_map[header] = i

    required = ["题型", "题干", "正确答案"]
    for h in required:
        if h not in col_map:
            raise ValidationError(f"缺少必需列: {h}")

    exam = None
    if exam_id:
        try:
            exam = Exam.objects.get(id=exam_id)
        except Exam.DoesNotExist:
            raise ValidationError(f"考试ID {exam_id} 不存在")

    success_count = 0
    errors: list[str] = []

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        try:
            if not row or not row[col_map["题干"]]:
                continue

            q_type_raw = str(row[col_map["题型"]] or "").strip()
            content = str(row[col_map["题干"]]).strip()
            answer = str(row[col_map["正确答案"]] or "").strip()
            score_val = row[col_map.get("分值", -1)] if "分值" in col_map else 10
            score = int(score_val) if score_val else 10
            difficulty_raw = str(row[col_map.get("难度", -1)] or "medium").strip() if "难度" in col_map else "medium"
            explanation = str(row[col_map.get("解析", -1)] or "").strip() if "解析" in col_map else ""

            # 解析题型
            q_type = QuestionType.CHOICE if q_type_raw in ("choice", "选择", "选择题") else QuestionType.BLANK

            # 解析选项
            options = []
            if q_type == QuestionType.CHOICE:
                for label in ["A", "B", "C", "D"]:
                    key = f"选项{label}"
                    if key in col_map and row[col_map[key]]:
                        options.append({"label": label, "text": str(row[col_map[key]]).strip()})

            # 解析难度
            difficulty_map = {"easy": Difficulty.EASY, "medium": Difficulty.MEDIUM, "hard": Difficulty.HARD}
            difficulty = difficulty_map.get(difficulty_raw, Difficulty.MEDIUM)

            question = Question.objects.create(
                exam=exam,
                question_type=q_type,
                content=content,
                options=json.dumps(options, ensure_ascii=False) if options else "[]",
                answer=answer,
                score=score,
                explanation=explanation,
                difficulty=difficulty,
                order=row_idx - 1,
            )

            if exam:
                ExamQuestion.objects.create(exam=exam, question=question, order=row_idx - 1)

            success_count += 1
        except Exception as e:
            errors.append(f"第{row_idx}行: {e}")

    wb.close()

    return {"success_count": success_count, "errors": errors}
