import logging
import tempfile
from typing import Any

from django.contrib.auth import authenticate
from django.core.cache import cache
from django.utils import timezone

from core.constants import LOCKOUT_MINUTES, MAX_LOGIN_ATTEMPTS, UserRole
from core.exceptions import (
    AuthenticationError,
    RateLimitError,
    ValidationError,
)
from core.utils import get_client_ip

from .models import LoginAttempt, User

logger = logging.getLogger(__name__)


class AuthService:
    """Authentication and user management service."""

    @staticmethod
    def authenticate_user(
        request,
        username: str,
        password: str,
        role: str | None = None,
    ) -> User:
        """Authenticate a user by username/password with rate limiting.

        Supports login by username, student_id, or email.
        Returns the authenticated User or raises an exception.
        """
        ip = get_client_ip(request)

        # Check IP rate limit
        if AuthService._is_ip_locked(ip):
            raise RateLimitError(f"登录尝试过多，请{LOCKOUT_MINUTES}分钟后再试")

        # Resolve login identifier to username
        resolved_username = AuthService._resolve_username(username)

        # Check if user exists but is inactive (authenticate() returns None for inactive users)
        try:
            target_user = User.objects.get(username=resolved_username)
            if not target_user.is_active:
                AuthService._record_attempt(ip, username, success=False)
                raise AuthenticationError("账号已被禁用")
        except User.DoesNotExist:
            pass

        # Authenticate
        user = authenticate(request, username=resolved_username, password=password)

        if user is None:
            AuthService._record_attempt(ip, username, success=False)
            raise AuthenticationError("用户名或密码错误")

        if role and user.role != role:
            AuthService._record_attempt(ip, username, success=False)
            raise AuthenticationError("用户角色不匹配")

        # Success — clear attempt counter and record
        AuthService._clear_attempts(ip)
        AuthService._record_attempt(ip, username, success=True)

        # Update last login IP
        user.last_login_ip = ip
        user.save(update_fields=["last_login_ip"])

        return user

    @staticmethod
    def change_password(user: User, old_password: str, new_password: str) -> None:
        """Change user password after verifying old password."""
        if not user.check_password(old_password):
            raise ValidationError("原密码错误")

        from core.validators import validate_password_complexity

        validate_password_complexity(new_password)

        user.set_password(new_password)
        user.must_change_password = False
        user.save(update_fields=["password", "must_change_password"])

    @staticmethod
    def import_students_from_excel(file) -> dict[str, Any]:
        """Import students from an Excel file.

        Returns: {"success_count": int, "errors": list[str]}
        """
        from core.validators import validate_file_extension, validate_file_size

        validate_file_extension(file)
        validate_file_size(file)

        try:
            from openpyxl import load_workbook

            wb = load_workbook(file, read_only=True)
            ws = wb.active
        except Exception as e:
            raise ValidationError(f"无法读取Excel文件: {e}")

        success_count = 0
        errors: list[str] = []

        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            try:
                if not row or not row[0]:
                    continue

                student_id = str(row[0]).strip()
                name = str(row[1]).strip() if len(row) > 1 and row[1] else student_id
                email = str(row[2]).strip() if len(row) > 2 and row[2] else ""

                if User.objects.filter(student_id=student_id).exists():
                    errors.append(f"第{row_idx}行: 学号{student_id}已存在")
                    continue

                # Default password: last 6 chars of student_id
                default_pwd = student_id[-6:] if len(student_id) >= 6 else student_id

                User.objects.create_user(
                    username=student_id,
                    password=default_pwd,
                    role=UserRole.STUDENT,
                    student_id=student_id,
                    first_name=name,
                    email=email,
                    must_change_password=True,
                )
                success_count += 1
            except Exception as e:
                errors.append(f"第{row_idx}行: {e}")

        wb.close()

        return {"success_count": success_count, "errors": errors}

    @staticmethod
    def _resolve_username(identifier: str) -> str:
        """Resolve student_id or email to username."""
        # Try student_id
        user = User.objects.filter(student_id=identifier).first()
        if user:
            return user.username
        # Try email
        user = User.objects.filter(email=identifier).first()
        if user:
            return user.username
        # Assume it's a username
        return identifier

    @staticmethod
    def _get_attempt_key(ip: str) -> str:
        return f"login_attempts:{ip}"

    @staticmethod
    def _is_ip_locked(ip: str) -> bool:
        key = AuthService._get_attempt_key(ip)
        attempts = cache.get(key, 0)
        return attempts >= MAX_LOGIN_ATTEMPTS

    @staticmethod
    def _record_attempt(ip: str, username: str, success: bool) -> None:
        LoginAttempt.objects.create(ip_address=ip, username=username, success=success)

        if not success:
            key = AuthService._get_attempt_key(ip)
            attempts = cache.get(key, 0) + 1
            cache.set(key, attempts, timeout=LOCKOUT_MINUTES * 60)

    @staticmethod
    def _clear_attempts(ip: str) -> None:
        key = AuthService._get_attempt_key(ip)
        cache.delete(key)
